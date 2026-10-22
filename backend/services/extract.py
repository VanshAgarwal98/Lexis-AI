"""
Text extraction from uploaded contract files (PDF / DOCX / TXT).
"""
import io
from pypdf import PdfReader
from docx import Document


def extract_text(file_bytes: bytes, filename: str) -> str:
    if not filename or "." not in filename:
        ext = ""
    else:
        ext = filename.lower().rsplit(".", 1)[-1]

    if ext == "pdf":
        return _extract_pdf(file_bytes)
    elif ext == "docx":
        return _extract_docx(file_bytes)
    elif ext in ("txt", "text", ""):
        return _extract_txt(file_bytes)
    else:
        raise ValueError(f"Unsupported file type: .{ext}. Please upload PDF, DOCX, or TXT.")


def _extract_pdf(file_bytes: bytes) -> str:
    reader = PdfReader(io.BytesIO(file_bytes))
    text_parts = []
    for page in reader.pages:
        try:
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)
        except Exception:
            continue
    return "\n\n".join(text_parts).strip()


def _extract_docx(file_bytes: bytes) -> str:
    doc = Document(io.BytesIO(file_bytes))
    text_parts = []
    
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt:
            text_parts.append(txt)
            
    for table in doc.tables:
        for row in table.rows:
            row_items = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_items:
                text_parts.append(" | ".join(row_items))
                
    return "\n\n".join(text_parts).strip()


def _extract_txt(file_bytes: bytes) -> str:
    for encoding in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
        try:
            return file_bytes.decode(encoding).strip()
        except UnicodeDecodeError:
            continue
    return file_bytes.decode("utf-8", errors="ignore").strip()
